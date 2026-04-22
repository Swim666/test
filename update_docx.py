#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
根据实验报告更新docx文档
"""

from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

# 读取实验报告内容
def read_experiment_report():
    with open('实验报告.md', 'r', encoding='utf-8') as f:
        content = f.read()
    return content

# 解析实验报告内容
def parse_report_content(content):
    sections = {}
    lines = content.strip().split('\n')
    
    current_section = None
    current_subsection = None
    section_content = []
    subsection_content = []
    
    for line in lines:
        line = line.strip()
        if line.startswith('# '):
            # 主标题
            if current_section:
                if current_subsection and subsection_content:
                    section_content.append((current_subsection, subsection_content))
                sections[current_section] = section_content
            current_section = line[2:]
            current_subsection = None
            section_content = []
            subsection_content = []
        elif line.startswith('## '):
            # 一级标题
            if current_subsection and subsection_content:
                section_content.append((current_subsection, subsection_content))
            current_subsection = line[3:]
            subsection_content = []
        elif line.startswith('### '):
            # 二级标题
            if current_subsection and subsection_content:
                section_content.append((current_subsection, subsection_content))
            current_subsection = line[4:]
            subsection_content = []
        elif line and current_subsection:
            # 内容行
            subsection_content.append(line)
        elif line and not current_subsection:
            # 主标题下的内容
            section_content.append(line)
    
    # 处理最后一个部分
    if current_subsection and subsection_content:
        section_content.append((current_subsection, subsection_content))
    if current_section:
        sections[current_section] = section_content
    
    return sections

# 创建docx文档
def create_docx(sections):
    doc = Document()
    
    # 添加标题
    main_title = 'Airbnb数据分析与建模实验报告'
    doc.add_heading(main_title, 0)
    
    # 处理各个章节
    sections_to_add = [
        ('一、实验目的', 1),
        ('二、实验数据集', 1),
        ('三、实验步骤', 1),
        ('四、实验结果与分析', 1),
        ('五、实验总结与改进', 1),
        ('六、附录', 1)
    ]
    
    for section_name, level in sections_to_add:
        section_content = sections.get(section_name, [])
        if section_content:
            doc.add_heading(section_name, level=level)
            for item in section_content:
                if isinstance(item, tuple):
                    # 子标题
                    subheading, content_lines = item
                    doc.add_heading(subheading, level=level+1)
                    for line in content_lines:
                        doc.add_paragraph(line)
                else:
                    # 直接内容
                    doc.add_paragraph(item)
    
    # 保存文档
    doc.save('Airbnb数据分析与建模_完整报告.docx')
    print('文档保存成功：Airbnb数据分析与建模_完整报告.docx')

if __name__ == '__main__':
    content = read_experiment_report()
    sections = parse_report_content(content)
    create_docx(sections)
